from __future__ import annotations

import json
import zipfile
from pathlib import Path
from typing import Any

import pandas as pd


def _clean(value: Any):
    """Convert pandas/numpy values to plain JSON-safe Python values."""
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    if hasattr(value, "item"):
        try:
            return value.item()
        except Exception:
            pass
    return value


def _records(df: pd.DataFrame) -> list[dict[str, Any]]:
    return [
        {str(k): _clean(v) for k, v in row.items()}
        for row in df.to_dict(orient="records")
    ]


def select_recommended(comparison: list[dict[str, Any]]) -> str | None:
    successful = [
        row
        for row in comparison
        if row.get("status") == "SUCCESS"
        and row.get("objective_score") is not None
    ]

    if not successful:
        return None

    def ranking_key(row):
        return (
            float(row.get("constraint_violations", 999999)),
            float(row.get("unassigned", 999999)),
            -float(row.get("assignment_rate", 0)),
            float(row.get("objective_score", float("inf"))),
            float(row.get("weight_imbalance_std", float("inf"))),
            float(row.get("runtime_seconds", float("inf"))),
        )

    best_result = min(successful, key=ranking_key)

    return best_result["algorithm"]

    if not successful:
        return None

    def number(row, *keys, default=0.0):
        for key in keys:
            value = row.get(key)
            if value is not None:
                try:
                    return float(value)
                except (TypeError, ValueError):
                    pass
        return float(default)

    def ranking_key(row):
        # 1. Prefer fewer constraint violations
        violations = number(
            row,
            "violations",
            "constraint_violations",
            default=999999,
        )

        # 2. Prefer more successfully assigned containers
        assigned = number(
            row,
            "assigned_containers",
            "assigned_count",
            "containers_assigned",
            default=0,
        )

        # If only assignment percentage exists, use that instead.
        assignment_rate = number(
            row,
            "assignment_rate",
            "assignment_percentage",
            default=0,
        )

        # 3. Lower overall objective penalty is better
        objective = number(
            row,
            "objective_score",
            default=float("inf"),
        )

        # 4. Lower weight imbalance is better
        imbalance = number(
            row,
            "imbalance",
            "weight_imbalance",
            default=float("inf"),
        )

        # 5. Runtime is only a final tie-breaker
        runtime = number(
            row,
            "runtime_seconds",
            "runtime",
            default=float("inf"),
        )

        return (
            violations,
            -assigned,
            -assignment_rate,
            objective,
            imbalance,
            runtime,
        )

    best_result = min(successful, key=ranking_key)

    return best_result["algorithm"]


def build_loading_plan(solution: pd.DataFrame) -> pd.DataFrame:
    """Create a presentation-friendly stowage/loading plan.

    Assigned containers are ordered by physical position: lower tiers first so the
    output can be read as a loading sequence. Unassigned containers are retained
    at the end so the report never hides demand that could not be stowed.
    """
    plan = solution.copy()
    if plan.empty:
        return plan

    assigned = plan[plan["slot_id"].notna()].copy()
    unassigned = plan[plan["slot_id"].isna()].copy()

    if not assigned.empty:
        sort_cols = [c for c in ["bay", "row", "tier", "destination_order", "priority"] if c in assigned.columns]
        assigned = assigned.sort_values(sort_cols, kind="stable").reset_index(drop=True)
        assigned.insert(0, "loading_order", range(1, len(assigned) + 1))
        assigned.insert(1, "assignment_status", "ASSIGNED")

    if not unassigned.empty:
        unassigned = unassigned.reset_index(drop=True)
        unassigned.insert(0, "loading_order", [None] * len(unassigned))
        unassigned.insert(1, "assignment_status", "UNASSIGNED")

    return pd.concat([assigned, unassigned], ignore_index=True)


def create_exports(
    result_dir: Path,
    run_id: str,
    dataset_name: str,
    scenario: dict[str, Any],
    comparison: list[dict[str, Any]],
    solutions: dict[str, pd.DataFrame],
    container_count: int,
    slot_count: int,
    total_runtime: float,
) -> tuple[dict[str, str], str | None, list[str]]:
    """Create CSV, JSON, XLSX, DOCX, PDF and ZIP exports for one run."""
    result_dir.mkdir(parents=True, exist_ok=True)
    errors: list[str] = []
    exports: dict[str, str] = {}

    recommended = select_recommended(comparison)
    recommended_solution = solutions.get(recommended, pd.DataFrame()).copy() if recommended else pd.DataFrame()
    recommended_plan = build_loading_plan(recommended_solution)

    base = result_dir / f"{run_id}_optimized_stowage"
    comparison_df = pd.DataFrame(comparison)

    # CSV: the actual recommended optimized plan, not only benchmark metrics.
    try:
        csv_path = Path(f"{base}.csv")
        recommended_plan.to_csv(csv_path, index=False)
        exports["csv"] = str(csv_path)
    except Exception as exc:
        errors.append(f"CSV export: {exc}")

    # JSON: run metadata + comparison + recommended plan.
    try:
        json_path = Path(f"{base}.json")
        payload = {
            "run_id": run_id,
            "dataset": dataset_name,
            "scenario": scenario,
            "container_count": container_count,
            "slot_count": slot_count,
            "total_runtime_seconds": total_runtime,
            "recommended_algorithm": recommended,
            "comparison": comparison,
            "recommended_plan": _records(recommended_plan),
        }
        json_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
        exports["json"] = str(json_path)
    except Exception as exc:
        errors.append(f"JSON export: {exc}")

    # Excel: summary, comparison, recommended plan and every successful solution.
    try:
        xlsx_path = Path(f"{base}.xlsx")
        with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
            summary_rows = [
                ["Run ID", run_id],
                ["Dataset", dataset_name],
                ["Scenario", scenario.get("scenario", "")],
                ["Route", scenario.get("route", "")],
                ["Dataset source", scenario.get("source", "")],
                ["Benchmark mode", scenario.get("benchmark_mode", "")],
                ["Source instance", scenario.get("source_instance", "")],
                ["Containers", container_count],
                ["Slots", slot_count],
                ["Demand / slot ratio", round(container_count / slot_count, 4) if slot_count else None],
                ["Recommended algorithm", recommended or "None"],
                ["Total runtime (s)", total_runtime],
            ]
            pd.DataFrame(summary_rows, columns=["Field", "Value"]).to_excel(writer, sheet_name="Summary", index=False)
            comparison_df.to_excel(writer, sheet_name="Algorithm Comparison", index=False)
            recommended_plan.to_excel(writer, sheet_name="Recommended Plan", index=False)
            for algorithm, solution in solutions.items():
                sheet = (algorithm[:25] + " Plan")[:31]
                build_loading_plan(solution).to_excel(writer, sheet_name=sheet, index=False)
        exports["xlsx"] = str(xlsx_path)
    except Exception as exc:
        errors.append(f"Excel export: {exc}")

    # Word report.
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.section import WD_ORIENT

        docx_path = Path(f"{base}.docx")
        doc = Document()
        section = doc.sections[0]
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section.top_margin = Inches(0.35)
        section.bottom_margin = Inches(0.35)
        section.left_margin = Inches(0.35)
        section.right_margin = Inches(0.35)

        title = doc.add_heading("Optimized Container Stowage Report", 0)
        title.runs[0].font.size = Pt(20)
        doc.add_paragraph(f"Run ID: {run_id}")
        doc.add_paragraph(f"Dataset: {dataset_name}")
        if scenario.get("scenario"):
            doc.add_paragraph(f"Scenario: {scenario['scenario']}")
        if scenario.get("route"):
            doc.add_paragraph(f"Route: {scenario['route']}")
        if scenario.get("source"):
            doc.add_paragraph(f"Dataset source: {scenario['source']}")
        if scenario.get("benchmark_mode"):
            doc.add_paragraph(f"Benchmark mode: {scenario['benchmark_mode']}")
        doc.add_paragraph(
            f"Demand: {container_count} containers | Capacity: {slot_count} slots | "
            f"Recommended algorithm: {recommended or 'None'}"
        )

        def set_table_font(table, size=7):
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(size)

        doc.add_heading("Algorithm Comparison", level=1)
        display_cols = [
            "algorithm", "status", "runtime_seconds", "assignment_rate",
            "slot_utilization", "constraint_violations", "rehandling", "objective_score",
        ]
        display_cols = [c for c in display_cols if c in comparison_df.columns]
        table = doc.add_table(rows=1, cols=len(display_cols))
        table.style = "Table Grid"
        for i, col in enumerate(display_cols):
            table.rows[0].cells[i].text = col.replace("_", " ").title()
        for _, row in comparison_df[display_cols].iterrows():
            cells = table.add_row().cells
            for i, col in enumerate(display_cols):
                value = _clean(row[col])
                cells[i].text = "" if value is None else str(value)

        set_table_font(table, 7)

        doc.add_heading("Recommended Stowage Plan", level=1)
        if recommended_plan.empty:
            doc.add_paragraph("No successful optimized plan was produced.")
        else:
            plan_cols = [
                "loading_order", "assignment_status", "container_id", "container_size",
                "container_weight", "destination", "priority", "slot_id", "bay", "row", "tier",
            ]
            plan_cols = [c for c in plan_cols if c in recommended_plan.columns]
            table = doc.add_table(rows=1, cols=len(plan_cols))
            table.style = "Table Grid"
            for i, col in enumerate(plan_cols):
                table.rows[0].cells[i].text = col.replace("_", " ").title()
            for _, row in recommended_plan[plan_cols].iterrows():
                cells = table.add_row().cells
                for i, col in enumerate(plan_cols):
                    value = _clean(row[col])
                    cells[i].text = "" if value is None else str(value)
            set_table_font(table, 7)

        doc.save(docx_path)
        exports["docx"] = str(docx_path)
    except Exception as exc:
        errors.append(f"Word export: {exc}")

    # PDF report.
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, PageBreak

        pdf_path = Path(f"{base}.pdf")
        styles = getSampleStyleSheet()
        story = [
            Paragraph("Optimized Container Stowage Report", styles["Title"]),
            Spacer(1, 4 * mm),
            Paragraph(f"<b>Run ID:</b> {run_id}", styles["BodyText"]),
            Paragraph(f"<b>Dataset:</b> {dataset_name}", styles["BodyText"]),
            Paragraph(f"<b>Scenario:</b> {scenario.get('scenario', '')}", styles["BodyText"]),
            Paragraph(f"<b>Route:</b> {scenario.get('route', '')}", styles["BodyText"]),
            Paragraph(f"<b>Dataset source:</b> {scenario.get('source', '')}", styles["BodyText"]),
            Paragraph(f"<b>Benchmark mode:</b> {scenario.get('benchmark_mode', '')}", styles["BodyText"]),
            Paragraph(
                f"<b>Demand:</b> {container_count} containers &nbsp;&nbsp; "
                f"<b>Capacity:</b> {slot_count} slots &nbsp;&nbsp; "
                f"<b>Recommended:</b> {recommended or 'None'}",
                styles["BodyText"],
            ),
            Spacer(1, 5 * mm),
            Paragraph("Algorithm Comparison", styles["Heading2"]),
        ]

        pdf_cols = [
            "algorithm", "status", "runtime_seconds", "assignment_rate",
            "slot_utilization", "constraint_violations", "rehandling", "objective_score",
        ]
        pdf_cols = [c for c in pdf_cols if c in comparison_df.columns]
        comp_data = [[c.replace("_", " ").title() for c in pdf_cols]]
        for _, row in comparison_df[pdf_cols].iterrows():
            comp_data.append(["" if _clean(row[c]) is None else str(_clean(row[c])) for c in pdf_cols])
        comp_table = LongTable(comp_data, repeatRows=1)
        comp_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF1FF")),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.grey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.extend([comp_table, PageBreak(), Paragraph("Recommended Stowage Plan", styles["Heading2"])])

        if recommended_plan.empty:
            story.append(Paragraph("No successful optimized plan was produced.", styles["BodyText"]))
        else:
            plan_cols = [
                "loading_order", "assignment_status", "container_id", "container_size",
                "container_weight", "destination", "priority", "slot_id", "bay", "row", "tier",
            ]
            plan_cols = [c for c in plan_cols if c in recommended_plan.columns]
            plan_data = [[c.replace("_", " ").title() for c in plan_cols]]
            for _, row in recommended_plan[plan_cols].iterrows():
                plan_data.append(["" if _clean(row[c]) is None else str(_clean(row[c])) for c in plan_cols])
            plan_table = LongTable(plan_data, repeatRows=1)
            plan_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF1FF")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#AAB3C2")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 6.2),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ]))
            story.append(plan_table)

        doc = SimpleDocTemplate(
            str(pdf_path),
            pagesize=landscape(A4),
            rightMargin=8 * mm,
            leftMargin=8 * mm,
            topMargin=8 * mm,
            bottomMargin=8 * mm,
        )
        doc.build(story)
        exports["pdf"] = str(pdf_path)
    except Exception as exc:
        errors.append(f"PDF export: {exc}")

    # One-click bundle.
    try:
        zip_path = Path(f"{base}_bundle.zip")
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for fmt, path in exports.items():
                if fmt == "zip":
                    continue
                p = Path(path)
                if p.exists():
                    zf.write(p, arcname=p.name)
        exports["zip"] = str(zip_path)
    except Exception as exc:
        errors.append(f"ZIP export: {exc}")

    return exports, recommended, errors
